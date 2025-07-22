import os
from typing import Awaitable, Callable
import cv2
import uuid
import pytest
import zipfile
import aiofiles
import mimetypes
import xlsxwriter
import numpy as np
from gtts import gTTS
from io import BytesIO
from docx import Document
from urllib.parse import quote
from reportlab.pdfgen import canvas
from reportlab.lib.pagesizes import letter
from PIL import Image, ImageDraw, ImageFont
from tests.http_client.AsyncHTTPClient import AsyncHTTPClient
from tests.constants import TEST_FILES_FOLDER
from pathlib import Path

FILES = "/files"
cwd = Path().cwd()

http_client = AsyncHTTPClient(timeout=1_000)


async def read_test_file(filename):
    async with aiofiles.open(filename, mode="rb") as f:
        return await f.read()


async def create_test_file(filetype: str, filename: str, text: str = "Test File"):
    match filetype:
        case "txt":
            await create_txt(filename=filename, text=text)

        case "pdf":
            await create_pdf(filename=filename, text=text)

        case "xls":
            await create_xls(filename=filename, text=text)

        case "docx":
            await create_docx(filename=filename, text=text)

        case "jpeg":
            await create_jpeg(filename=filename, text=text)

        case "png":
            await create_png(filename=filename, text=text)

        case "zip":
            await create_zip(filename=filename, text=text)

        case "html":
            await create_html(filename=filename, text=text)

        case "mp3":
            await create_mp3(filename=filename, text=text)

        case "mp4":
            await create_mp4(filename=filename, text=text)

        case _:
            assert False, f"Invalid file type: {filetype}"


async def create_txt(filename: str, text: str):
    async with aiofiles.open(filename, mode="w") as file:
        await file.write(text)

    mime_type, _ = mimetypes.guess_type(filename)

    assert mime_type == "text/plain"


async def create_pdf(filename: str, text: str):
    c = canvas.Canvas(filename, pagesize=letter)

    _, height = letter

    c.drawString(100, height - 100, text)

    c.save()

    async with aiofiles.open(filename, "wb") as f:
        await f.write(c.getpdfdata())

    mime_type, _ = mimetypes.guess_type(filename)

    assert mime_type == "application/pdf"


async def create_xls(filename: str, text: str):
    output = BytesIO()

    workbook = xlsxwriter.Workbook(output, {"in_memory": True})
    worksheet = workbook.add_worksheet()

    worksheet.write("A1", text)

    workbook.close()

    xls_data = output.getvalue()

    async with aiofiles.open(filename, "wb") as f:
        await f.write(xls_data)

    mime_type, _ = mimetypes.guess_type(filename)

    assert mime_type == "application/vnd.ms-excel"


async def create_docx(filename: str, text: str):
    output = BytesIO()

    doc = Document()

    doc.add_paragraph(text)

    doc.save(output)

    output.seek(0)

    docx_data = output.getvalue()

    async with aiofiles.open(filename, "wb") as f:
        await f.write(docx_data)

    mime_type, _ = mimetypes.guess_type(filename)

    assert (
        mime_type
        == "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    )


async def create_jpeg(filename: str, text: str):
    image = Image.new("RGB", (200, 100), color="white")

    draw = ImageDraw.Draw(image)

    draw.text((10, 50), text, fill="black")

    output = BytesIO()

    image.save(output, format="JPEG")

    output.seek(0)

    async with aiofiles.open(filename, "wb") as f:
        await f.write(output.getvalue())

    mime_type, _ = mimetypes.guess_type(filename)

    assert mime_type == "image/jpeg"


async def create_png(filename: str, text: str):
    image = Image.new("RGB", (200, 100), color="white")

    draw = ImageDraw.Draw(image)

    draw.text((10, 50), text, fill="black")

    output = BytesIO()

    image.save(output, format="PNG")

    output.seek(0)

    async with aiofiles.open(filename, "wb") as f:
        await f.write(output.getvalue())

    mime_type, _ = mimetypes.guess_type(filename)

    assert mime_type == "image/png"


async def create_zip(filename: str, text: str):
    txt_file = f"{cwd}{TEST_FILES_FOLDER}/test.txt"

    async with aiofiles.open(f"{cwd}{TEST_FILES_FOLDER}/test.txt", "w") as file:
        await file.write(text)

    with zipfile.ZipFile(filename, "w") as zipf:
        zipf.write(txt_file)

    mime_type, _ = mimetypes.guess_type(filename)

    assert mime_type == "application/zip"


async def create_html(filename: str, text: str):
    html_content = """
    <!DOCTYPE html>
    <html lang="en">
    <head>
        <meta charset="UTF-8">
        <meta name="viewport" content="width=device-width, initial-scale=1.0">
        <title>Test HTML</title>
    </head>
    <body>
        <h1>{h1}</h1>
        <p>{p}</p>
    </body>
    </html>
    """

    html_content = html_content.format(h1=text, p=text)

    async with aiofiles.open(filename, "w") as file:
        await file.write(html_content)

    mime_type, _ = mimetypes.guess_type(filename)

    assert mime_type == "text/html"


async def create_mp3(filename: str, text: str):
    tts = gTTS(text=text, lang="en")

    temp_file = "temp.mp3"

    tts.save(temp_file)

    async with aiofiles.open(temp_file, mode="rb") as f_in:
        async with aiofiles.open(filename, mode="wb") as f_out:
            await f_out.write(await f_in.read())

    os.remove(temp_file)

    mime_type, _ = mimetypes.guess_type(filename)

    assert mime_type == "audio/mpeg"


async def create_mp4(filename: str, text: str):
    duration: int = 5

    fps: int = 24

    width, height = 640, 480

    background_color = (0, 0, 0)

    image = Image.new("RGB", (width, height), background_color)

    draw = ImageDraw.Draw(image)

    font = ImageFont.load_default()

    bbox = draw.textbbox((0, 0), text, font=font)

    text_width = bbox[2] - bbox[0]

    text_height = bbox[3] - bbox[1]

    text_position = ((width - text_width) // 2, (height - text_height) // 2)

    draw.text(text_position, text, font=font, fill=(255, 255, 255))

    frame = np.array(image)

    frame = cv2.cvtColor(frame, cv2.COLOR_RGB2BGR)

    fourcc = cv2.VideoWriter_fourcc(*"mp4v")

    temp_file = "temp.mp4"

    video_writer = cv2.VideoWriter(temp_file, fourcc, fps, (width, height))

    for _ in range(duration * fps):
        video_writer.write(frame)

    video_writer.release()

    async with aiofiles.open(temp_file, mode="rb") as f_in:
        async with aiofiles.open(filename, mode="wb") as f_out:
            await f_out.write(await f_in.read())

    os.remove(temp_file)

    mime_type, _ = mimetypes.guess_type(filename)

    assert mime_type == "video/mp4"


@pytest.mark.asyncio
@pytest.mark.parametrize(
    "file_extension",
    [
        ("txt"),
        ("pdf"),
        ("xls"),
        ("docx"),
        ("jpeg"),
        ("png"),
        ("zip"),
        ("html"),
        ("mp3"),
        ("mp4"),
    ],
    ids=[
        "upload txt file",
        "upload pdf file",
        "upload xls file",
        "upload docx file",
        "upload jpeg file",
        "upload png file",
        "upload zip file",
        "upload html file",
        "upload mp3 file",
        "upload mp4 file",
    ],
)
async def test_files_valid_post(file_extension, user_jwt_token: str):
    filename = f"{cwd}{TEST_FILES_FOLDER}/test.{file_extension}"

    await create_test_file(filetype=file_extension, filename=filename)

    mime_type, _ = mimetypes.guess_type(filename)

    file_id = await http_client.upload_file(
        path=FILES,
        filename=filename,
        request_id=str(uuid.uuid4()),
        session_id=str(uuid.uuid4()),
        content_type=mime_type,
        headers={"Authorization": f"Bearer {user_jwt_token}"},
    )

    assert file_id, "Invalid file id received"


@pytest.mark.asyncio
@pytest.mark.parametrize(
    "file_extension, content_type",
    [
        ("txt", None),
        ("txt", "application/pdf"),
    ],
    ids=[
        "upload txt file without content type",
        "upload txt file with wrong content type",
    ],
)
async def test_files_content_type_post(
    file_extension, content_type, user_jwt_token: str
):
    filename = f"{cwd}{TEST_FILES_FOLDER}/test.{file_extension}"

    await create_test_file(filetype=file_extension, filename=filename)

    file_id = await http_client.upload_file(
        path=FILES,
        filename=filename,
        request_id=str(uuid.uuid4()),
        session_id=str(uuid.uuid4()),
        content_type=content_type,
        headers={"Authorization": f"Bearer {user_jwt_token}"},
    )

    assert file_id, "Invalid file id received"


@pytest.mark.asyncio
@pytest.mark.parametrize(
    "file_extension, request_id, session_id",
    [
        ("txt", None, None),
        ("txt", str(uuid.uuid4()), None),
        ("txt", None, str(uuid.uuid4())),
        ("txt", str(uuid.uuid4())[:-1], str(uuid.uuid4())),
        ("txt", str(uuid.uuid4()), str(uuid.uuid4())[:-1]),
    ],
    ids=[
        "upload txt file with None as request_id and session_id",
        "upload txt file with None as session_id",
        "upload txt file with None as request_id",
        "upload txt file with invalid request_id",
        "upload txt file with invalid session_id",
    ],
)
async def test_files_request_session_ids_post(
    file_extension, request_id, session_id, user_jwt_token: str
):
    filename = f"{cwd}{TEST_FILES_FOLDER}/test.{file_extension}"

    await create_test_file(filetype=file_extension, filename=filename)

    mime_type, _ = mimetypes.guess_type(filename)

    await http_client.upload_file(
        path=FILES,
        filename=filename,
        request_id=request_id,
        session_id=session_id,
        content_type=mime_type,
        expected_status_codes=[400],
        headers={"Authorization": f"Bearer {user_jwt_token}"},
    )


@pytest.mark.asyncio
@pytest.mark.parametrize(
    "file_extension, text",
    [
        ("txt", "Test Example"),
    ],
    ids=[
        "upload txt file and get its content",
    ],
)
async def test_files_txt_file_id_get(file_extension, text, user_jwt_token: str):
    filename = f"{cwd}{TEST_FILES_FOLDER}/test.{file_extension}"

    await create_test_file(filetype=file_extension, filename=filename, text=text)

    mime_type, _ = mimetypes.guess_type(filename)

    file_id = await http_client.upload_file(
        path=FILES,
        filename=filename,
        request_id=str(uuid.uuid4()),
        session_id=str(uuid.uuid4()),
        content_type=mime_type,
        headers={"Authorization": f"Bearer {user_jwt_token}"},
    )

    assert file_id

    file_content = await http_client.get(
        path=f"{FILES}/{file_id}",
        headers={"Authorization": f"Bearer {user_jwt_token}"},
    )

    assert file_content == text, f"Invalid file content: {file_content}"


@pytest.mark.asyncio
@pytest.mark.parametrize(
    "file_extension, text",
    [
        ("pdf", "Test Example"),
        ("xls", "Test Example"),
        ("docx", "Test Example"),
        ("jpeg", "Test Example"),
        ("png", "Test Example"),
        ("zip", "Test Example"),
        ("html", "Test Example"),
        ("mp3", "Test Example"),
        ("mp4", "Test Example"),
    ],
    ids=[
        "upload pdf file and get its content",
        "upload xls file and get its content",
        "upload docx file and get its content",
        "upload jpeg file and get its content",
        "upload png file and get its content",
        "upload zip file and get its content",
        "upload html file and get its content",
        "upload mp3 file and get its content",
        "upload mp4 file and get its content",
    ],
)
async def test_files_file_id_get(file_extension, text, user_jwt_token: str):
    filename = f"{cwd}{TEST_FILES_FOLDER}/test.{file_extension}"

    await create_test_file(filetype=file_extension, filename=filename, text=text)

    mime_type, _ = mimetypes.guess_type(filename)

    file_id = await http_client.upload_file(
        path=FILES,
        filename=filename,
        request_id=str(uuid.uuid4()),
        session_id=str(uuid.uuid4()),
        content_type=mime_type,
        headers={"Authorization": f"Bearer {user_jwt_token}"},
    )

    assert file_id

    file_in_bytest = await read_test_file(filename=filename)

    FILE_ID_CONTENT = f"{FILES}/{file_id}"

    file_content = await http_client.get(
        path=FILE_ID_CONTENT,
        headers={"Authorization": f"Bearer {user_jwt_token}"},
    )

    assert file_content == file_in_bytest, f"Invalid file content: {file_content}"


@pytest.mark.asyncio
@pytest.mark.parametrize(
    "file_extension, text, session_id, request_id",
    [
        ("txt", "Test Example", str(uuid.uuid4()), str(uuid.uuid4())),
        ("pdf", "Test Example", str(uuid.uuid4()), str(uuid.uuid4())),
        ("xls", "Test Example", str(uuid.uuid4()), str(uuid.uuid4())),
        ("docx", "Test Example", str(uuid.uuid4()), str(uuid.uuid4())),
        ("jpeg", "Test Example", str(uuid.uuid4()), str(uuid.uuid4())),
        ("png", "Test Example", str(uuid.uuid4()), str(uuid.uuid4())),
        ("zip", "Test Example", str(uuid.uuid4()), str(uuid.uuid4())),
        ("html", "Test Example", str(uuid.uuid4()), str(uuid.uuid4())),
        ("mp3", "Test Example", str(uuid.uuid4()), str(uuid.uuid4())),
        ("mp4", "Test Example", str(uuid.uuid4()), str(uuid.uuid4())),
    ],
    ids=[
        "upload txt file and get its metadata",
        "upload pdf file and get its metadata",
        "upload xls file and get its metadata",
        "upload docx file and get its metadata",
        "upload jpeg file and get its metadata",
        "upload png file and get its metadata",
        "upload zip file and get its metadata",
        "upload html file and get its metadata",
        "upload mp3 file and get its metadata",
        "upload mp4 file and get its metadata",
    ],
)
async def test_files_file_id_metadata_get(
    file_extension,
    text,
    session_id,
    request_id,
    user_jwt_token: str,
    get_user: Callable[[str], Awaitable[str]],
):
    filename = f"{cwd}{TEST_FILES_FOLDER}/test.{file_extension}"

    await create_test_file(filetype=file_extension, filename=filename, text=text)

    mime_type, _ = mimetypes.guess_type(filename)

    original_name = quote(filename, safe="")

    file_id = await http_client.upload_file(
        path=FILES,
        filename=filename,
        request_id=session_id,
        session_id=request_id,
        content_type=mime_type,
        headers={"Authorization": f"Bearer {user_jwt_token}"},
    )

    assert file_id

    FILE_ID_METADATA = f"{FILES}/{file_id}/metadata"
    file_metadata = await http_client.get(
        path=FILE_ID_METADATA,
        headers={"Authorization": f"Bearer {user_jwt_token}"},
    )

    expected_file_metadata = {
        "id": file_id,
        # "session_id": session_id,
        # "request_id": request_id,
        "session_id": None,
        "request_id": None,
        "original_name": original_name,
        "mimetype": mime_type,
        "internal_id": file_id,
        "internal_name": f"{file_id}.{file_extension}",
        "from_agent": False,
        "creator_id": await get_user(user_jwt_token),
    }

    assert file_metadata == expected_file_metadata, "Received invalid metadata"
