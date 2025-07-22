import os
from typing import Awaitable, Callable
import cv2
import json
import uuid
import pytest
import asyncio
import zipfile
import aiofiles
import mimetypes
import websockets
import xlsxwriter
import numpy as np

from gtts import gTTS
from io import BytesIO
from loguru import logger
from copy import deepcopy
from docx import Document
from urllib.parse import quote
from reportlab.pdfgen import canvas
from reportlab.lib.pagesizes import letter
from PIL import Image, ImageDraw, ImageFont
from genai_session.session import GenAISession
from genai_session.utils.context import GenAIContext
from tests.conftest import DummyAgent
from tests.http_client.AsyncHTTPClient import AsyncHTTPClient
from tests.constants import URI, TEST_FILES_FOLDER, WS_HEADERS, WS_MESSAGE
from pathlib import Path

cwd = Path().cwd()
mimetypes.init()


http_client = AsyncHTTPClient(timeout=10)

FILES_ENDPOINT = "/files"

# AUTH_JWT_FILE = str(uuid.uuid4())
# AUTH_JWT_METADATA = str(uuid.uuid4())
# AUTH_JWT_SAVER = str(uuid.uuid4())

# file_session = GenAISession(jwt_token=AUTH_JWT_FILE)
# metadata_session = GenAISession(jwt_token=AUTH_JWT_METADATA)
# file_saver_session = GenAISession(jwt_token=AUTH_JWT_SAVER)


async def read_test_file(filename):
    async with aiofiles.open(filename, mode="rb") as f:
        return await f.read()


async def create_test_file(filetype: str, filename: str, text: str = "Test File"):
    match filetype:
        case "txt":
            await create_txt(filename=filename, text=text)

        case "pdf":
            await create_pdf(filename=filename, text=text)

        case "xls":
            await create_xls(filename=filename, text=text)

        case "docx":
            await create_docx(filename=filename, text=text)

        case "jpeg":
            await create_jpeg(filename=filename, text=text)

        case "png":
            await create_png(filename=filename, text=text)

        case "zip":
            await create_zip(filename=filename, text=text)

        case "html":
            await create_html(filename=filename, text=text)

        case "mp3":
            await create_mp3(filename=filename, text=text)

        case "mp4":
            await create_mp4(filename=filename, text=text)

        case _:
            assert False, f"Invalid file type: {filetype}"


async def create_txt(filename: str, text: str):
    async with aiofiles.open(filename, mode="w") as file:
        await file.write(text)

    mime_type, _ = mimetypes.guess_type(filename)

    assert mime_type == "text/plain"


async def create_pdf(filename: str, text: str):
    c = canvas.Canvas(filename, pagesize=letter)

    _, height = letter

    c.drawString(100, height - 100, text)

    c.save()

    async with aiofiles.open(filename, "wb") as f:
        await f.write(c.getpdfdata())

    mime_type, _ = mimetypes.guess_type(filename)

    assert mime_type == "application/pdf"


async def create_xls(filename: str, text: str):
    output = BytesIO()

    workbook = xlsxwriter.Workbook(output, {"in_memory": True})
    worksheet = workbook.add_worksheet()

    worksheet.write("A1", text)

    workbook.close()

    xls_data = output.getvalue()

    async with aiofiles.open(filename, "wb") as f:
        await f.write(xls_data)

    mime_type, _ = mimetypes.guess_type(filename)

    assert mime_type == "application/vnd.ms-excel"


async def create_docx(filename: str, text: str):
    output = BytesIO()

    doc = Document()

    doc.add_paragraph(text)

    doc.save(output)

    output.seek(0)

    docx_data = output.getvalue()

    async with aiofiles.open(filename, "wb") as f:
        await f.write(docx_data)

    mime_type, _ = mimetypes.guess_type(filename)

    assert (
        mime_type
        == "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    )


async def create_jpeg(filename: str, text: str):
    image = Image.new("RGB", (200, 100), color="white")

    draw = ImageDraw.Draw(image)

    draw.text((10, 50), text, fill="black")

    output = BytesIO()

    image.save(output, format="JPEG")

    output.seek(0)

    async with aiofiles.open(filename, "wb") as f:
        await f.write(output.getvalue())

    mime_type, _ = mimetypes.guess_type(filename)

    assert mime_type == "image/jpeg"


async def create_png(filename: str, text: str):
    image = Image.new("RGB", (200, 100), color="white")

    draw = ImageDraw.Draw(image)

    draw.text((10, 50), text, fill="black")

    output = BytesIO()

    image.save(output, format="PNG")

    output.seek(0)

    async with aiofiles.open(filename, "wb") as f:
        await f.write(output.getvalue())

    mime_type, _ = mimetypes.guess_type(filename)

    assert mime_type == "image/png"


async def create_zip(filename: str, text: str):
    txt_file = f"{cwd}{TEST_FILES_FOLDER}/test.txt"

    async with aiofiles.open(f"{cwd}{TEST_FILES_FOLDER}/test.txt", "w") as file:
        await file.write(text)

    with zipfile.ZipFile(filename, "w") as zipf:
        zipf.write(txt_file)

    mime_type, _ = mimetypes.guess_type(filename)

    assert mime_type == "application/zip"


async def create_html(filename: str, text: str):
    html_content = """
    <!DOCTYPE html>
    <html lang="en">
    <head>
        <meta charset="UTF-8">
        <meta name="viewport" content="width=device-width, initial-scale=1.0">
        <title>Test HTML</title>
    </head>
    <body>
        <h1>{h1}</h1>
        <p>{p}</p>
    </body>
    </html>
    """

    html_content = html_content.format(h1=text, p=text)

    async with aiofiles.open(filename, "w") as file:
        await file.write(html_content)

    mime_type, _ = mimetypes.guess_type(filename)

    assert mime_type == "text/html"


async def create_mp3(filename: str, text: str):
    tts = gTTS(text=text, lang="en")

    temp_file = "temp.mp3"

    tts.save(temp_file)

    async with aiofiles.open(temp_file, mode="rb") as f_in:
        async with aiofiles.open(filename, mode="wb") as f_out:
            await f_out.write(await f_in.read())

    os.remove(temp_file)

    mime_type, _ = mimetypes.guess_type(filename)

    assert mime_type == "audio/mpeg"


async def create_mp4(filename: str, text: str):
    duration: int = 5

    fps: int = 24

    width, height = 640, 480

    background_color = (0, 0, 0)

    image = Image.new("RGB", (width, height), background_color)

    draw = ImageDraw.Draw(image)

    font = ImageFont.load_default()

    bbox = draw.textbbox((0, 0), text, font=font)

    text_width = bbox[2] - bbox[0]

    text_height = bbox[3] - bbox[1]

    text_position = ((width - text_width) // 2, (height - text_height) // 2)

    draw.text(text_position, text, font=font, fill=(255, 255, 255))

    frame = np.array(image)

    frame = cv2.cvtColor(frame, cv2.COLOR_RGB2BGR)

    fourcc = cv2.VideoWriter_fourcc(*"mp4v")

    temp_file = "temp.mp4"

    video_writer = cv2.VideoWriter(temp_file, fourcc, fps, (width, height))

    for _ in range(duration * fps):
        video_writer.write(frame)

    video_writer.release()

    async with aiofiles.open(temp_file, mode="rb") as f_in:
        async with aiofiles.open(filename, mode="wb") as f_out:
            await f_out.write(await f_in.read())

    os.remove(temp_file)

    mime_type, _ = mimetypes.guess_type(filename)

    assert mime_type == "video/mp4"


# @file_session.bind(name="File data getter", description="Gets uploaded file bytes")
# async def get_file_by_id(agent_context: GenAIContext, file_id: str):
#     file = await agent_context.files.get_by_id(file_id)

#     return file.getvalue().hex()


# @file_saver_session.bind(name="File data saver", description="Saves files data")
# async def save_file_content(agent_context: GenAIContext, content: str, filename: str):
#     bytes_content = bytes.fromhex(content)

#     file_id = await agent_context.files.save(content=bytes_content, filename=filename)

#     return file_id


# @metadata_session.bind(
#     name="File metadata getter", description="Gets metadata of uploaded file"
# )
# async def get_file_metadata_by_id(agent_context: GenAIContext, file_id: str):
#     metadata = await agent_context.files.get_metadata_by_id(file_id)

#     return metadata


# async def start_file_getter():
#     logger.info(f"Agent with ID {AUTH_JWT_FILE} started")

#     await file_session.process_events()


# async def start_file_saver():
#     logger.info(f"Agent with ID {AUTH_JWT_SAVER} started")

#     await file_saver_session.process_events()


# async def start_file_metadata_getter():
#     logger.info(f"Agent with ID {AUTH_JWT_METADATA} started")

#     await metadata_session.process_events()


@pytest.mark.asyncio
@pytest.mark.parametrize(
    "request_id, session_id, file_extension, mimetype",
    [
        (None, None, ".txt", "text/plain"),
        (str(uuid.uuid4()), str(uuid.uuid4()), ".txt", "text/plain"),
        (str(uuid.uuid4()), None, ".txt", "text/plain"),
        (None, str(uuid.uuid4()), ".txt", "text/plain"),
    ],
    ids=[
        "get file metadata without request_id and session_id",
        "get file metadata with request_id and session_id",
        "get file metadata with request_id and no session_id",
        "get file metadata with no request_id and with session_id",
    ],
)
async def test_agent_get_file_metadata(
    request_id,
    session_id,
    file_extension,
    mimetype,
    user_jwt_token: str,
    agent_jwt_factory: Callable[[str], Awaitable[str]],
    get_user: Callable[[str], Awaitable[str]],
):
    metadata_jwt_token = await agent_jwt_factory(user_jwt_token)
    metadata_session = GenAISession(jwt_token=metadata_jwt_token)

    @metadata_session.bind(
        name="File metadata getter", description="Gets metadata of uploaded file"
    )
    async def get_file_metadata_by_id(agent_context: GenAIContext, file_id: str):
        metadata = await agent_context.files.get_metadata_by_id(file_id)

        return metadata

    async def start_file_metadata_getter():
        await metadata_session.process_events()

    txt_file = f"test{file_extension}"

    filename = f"{cwd}{TEST_FILES_FOLDER}/{txt_file}"

    content = "Test File"

    original_name = quote(filename, safe="")

    await create_txt(filename=filename, text=content)

    mime_type, _ = mimetypes.guess_type(filename)

    file_id = await http_client.upload_file(
        path=FILES_ENDPOINT,
        filename=filename,
        request_id=request_id,
        session_id=session_id,
        content_type=mime_type,
        headers={"Authorization": f"Bearer {user_jwt_token}"},
    )

    event_task_1 = asyncio.create_task(start_file_metadata_getter())

    await asyncio.sleep(0.1)

    event_tasks = [event_task_1]

    message = deepcopy(WS_MESSAGE)

    message["agent_uuid"] = metadata_session.agent_id

    message["request_payload"] = {"file_id": file_id}

    try:
        async with websockets.connect(URI, additional_headers=WS_HEADERS) as websocket:
            await websocket.send(json.dumps(message))

            response = await websocket.recv()

            response_data = json.loads(response)

            logger.info(f"WebSocket Response: {response_data}")

            del response_data["execution_time"]

            expected_response_data = {
                "response": {
                    "id": file_id,
                    "session_id": None,
                    "request_id": None,
                    "original_name": original_name,
                    "mimetype": mimetype,
                    "internal_id": file_id,
                    "internal_name": f"{file_id}{file_extension}",
                    "from_agent": False,
                    "creator_id": await get_user(user_jwt_token),
                },
                "message_type": "agent_response",
            }

            assert response_data == expected_response_data

            logger.info("Test successful!")

    finally:
        for task in event_tasks:
            task.cancel()

        try:
            await task

        except asyncio.CancelledError:
            logger.info("Background task has been properly cancelled.")


@pytest.mark.asyncio
@pytest.mark.parametrize(
    "agent_uuid, file_extension",
    [
        (str(uuid.uuid4()), ".txt"),
        (None, ".txt"),
    ],
    ids=[
        "get file metadata with invalid agent_id",
        "get file metadata with agent_id as None",
    ],
)
async def test_agent_get_file_invalid_agent_id(
    agent_uuid,
    file_extension,
    user_jwt_token: str,
    agent_jwt_factory: Callable[[str], Awaitable[str]],
):
    metadata_jwt_token = await agent_jwt_factory(user_jwt_token)
    metadata_session = GenAISession(jwt_token=metadata_jwt_token)

    @metadata_session.bind(
        name="File metadata getter", description="Gets metadata of uploaded file"
    )
    async def get_file_metadata_by_id(agent_context: GenAIContext, file_id: str):
        metadata = await agent_context.files.get_metadata_by_id(file_id)

        return metadata

    async def start_file_metadata_getter():
        await metadata_session.process_events()

    txt_file = f"test{file_extension}"

    filename = f"{cwd}{TEST_FILES_FOLDER}/{txt_file}"

    content = "Test File"

    await create_txt(filename=filename, text=content)

    file_id = await http_client.upload_file(
        path=FILES_ENDPOINT,
        filename=filename,
        headers={"Authorization": f"Bearer {user_jwt_token}"},
    )

    event_task_1 = asyncio.create_task(start_file_metadata_getter())

    await asyncio.sleep(0.1)

    event_tasks = [event_task_1]

    message = deepcopy(WS_MESSAGE)

    message["agent_uuid"] = agent_uuid

    message["request_payload"] = {"file_id": file_id}

    try:
        async with websockets.connect(URI, additional_headers=WS_HEADERS) as websocket:
            await websocket.send(json.dumps(message))

            response = await websocket.recv()

            response_data = json.loads(response)

            logger.info(f"WebSocket Response: {response_data}")

            expected_response_data = {
                "message_type": "agent_error",
                "error": {
                    "error_message": "Agent is NOT active",
                    "error_type": "AgentNotActive",
                },
            }

            assert response_data == expected_response_data

            logger.info("Test successful!")

    finally:
        for task in event_tasks:
            task.cancel()

        try:
            await task

        except asyncio.CancelledError:
            logger.info("Background task has been properly cancelled.")


@pytest.mark.asyncio
@pytest.mark.parametrize(
    "file_id, reason",
    [
        (str(uuid.uuid4()), "Failed to retrieve file: 400, message='Bad Request'"),
        (None, "Failed to retrieve file: 422, message='Unprocessable Entity'"),
    ],
    ids=[
        "get file metadata with invalid file_id",
        "get file metadata with file_id as None",
    ],
)
async def test_agent_get_file_invalid_file_id(
    file_id,
    reason,
    user_jwt_token: str,
    agent_jwt_factory: Callable[[str], Awaitable[str]],
    agent_factory: Callable[[], DummyAgent],
):
    metadata_jwt_token = await agent_jwt_factory(user_jwt_token)
    metadata_session = GenAISession(jwt_token=metadata_jwt_token)

    @metadata_session.bind(
        name="File metadata getter", description="Gets metadata of uploaded file"
    )
    async def get_file_metadata_by_id(agent_context: GenAIContext, file_id: str):
        metadata = await agent_context.files.get_metadata_by_id(file_id)

        return metadata

    async def start_file_metadata_getter():
        await metadata_session.process_events()

    txt_file = "test.txt"

    filename = f"{cwd}{TEST_FILES_FOLDER}/{txt_file}"

    content = "Test File"

    await create_txt(filename=filename, text=content)

    _ = await http_client.upload_file(path=FILES_ENDPOINT, filename=filename)

    event_task_1 = asyncio.create_task(start_file_metadata_getter())

    await asyncio.sleep(0.1)

    event_tasks = [event_task_1]

    message = deepcopy(WS_MESSAGE)

    message["agent_uuid"] = metadata_session.agent_id

    message["request_payload"] = {"file_id": file_id}

    try:
        async with websockets.connect(URI, additional_headers=WS_HEADERS) as websocket:
            await websocket.send(json.dumps(message))

            response = await websocket.recv()

            response_data = json.loads(response)

            logger.info(f"WebSocket Response: {response_data}")

            expected_response_data = {
                "error": {
                    "error_message": f"{reason}, url='http://localhost:8000/files/{file_id}/metadata'"
                },
                "execution_time": 0,
                "message_type": "agent_error",
            }

            assert response_data == expected_response_data

            logger.info("Test successful!")

    finally:
        for task in event_tasks:
            task.cancel()

        try:
            await task

        except asyncio.CancelledError:
            logger.info("Background task has been properly cancelled.")


@pytest.mark.asyncio
@pytest.mark.parametrize(
    "request_id, session_id, file_extension, mimetype",
    [
        (str(uuid.uuid4()), str(uuid.uuid4()), "pdf", "application/pdf"),
        (str(uuid.uuid4()), str(uuid.uuid4()), "txt", "text/plain"),
        (str(uuid.uuid4()), str(uuid.uuid4()), "xls", "application/vnd.ms-excel"),
        (
            str(uuid.uuid4()),
            str(uuid.uuid4()),
            "docx",
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        ),
        (str(uuid.uuid4()), str(uuid.uuid4()), "jpeg", "image/jpeg"),
        (str(uuid.uuid4()), str(uuid.uuid4()), "png", "image/png"),
        (str(uuid.uuid4()), str(uuid.uuid4()), "zip", "application/zip"),
        (str(uuid.uuid4()), str(uuid.uuid4()), "html", "text/html"),
        (str(uuid.uuid4()), str(uuid.uuid4()), "mp3", "audio/mpeg"),
        (str(uuid.uuid4()), str(uuid.uuid4()), "mp4", "video/mp4"),
    ],
    ids=[
        "save pdf file and get its metadata",
        "save txt file and get its metadata",
        "save xls file and get its metadata",
        "save docs file and get its metadata",
        "save jpeg file and get its metadata",
        "save png file and get its metadata",
        "save zip file and get its metadata",
        "save html file and get its metadata",
        "save mp3 file and get its metadata",
        "save mp4 file and get its metadata",
    ],
)
async def test_agent_save_file(
    request_id,
    session_id,
    file_extension,
    mimetype,
    user_jwt_token: str,
    agent_jwt_factory: Callable[[str], Awaitable[str]],
    agent_factory: Callable[[], DummyAgent],
    get_user: Callable[[], DummyAgent],
):
    # metadata -----------------------------
    metadata_jwt_token = await agent_jwt_factory(user_jwt_token)
    metadata_session = GenAISession(jwt_token=metadata_jwt_token)

    @metadata_session.bind(
        name="File metadata getter", description="Gets metadata of uploaded file"
    )
    async def get_file_metadata_by_id(agent_context: GenAIContext, file_id: str):
        metadata = await agent_context.files.get_metadata_by_id(file_id)

        return metadata

    async def start_file_metadata_getter():
        await metadata_session.process_events()

    # ------------------------------------------------------
    # File saver
    filesaver_jwt_token = await agent_jwt_factory(user_jwt_token)
    file_saver_session = GenAISession(jwt_token=filesaver_jwt_token)

    @file_saver_session.bind(name="File data saver", description="Saves files data")
    async def save_file_content(
        agent_context: GenAIContext, content: str, filename: str
    ):
        bytes_content = bytes.fromhex(content)

        file_id = await agent_context.files.save(
            content=bytes_content, filename=filename
        )

        return file_id

    async def start_file_saver():
        await file_saver_session.process_events()

    file = f"test.{file_extension}"

    filename = f"{cwd}{TEST_FILES_FOLDER}/{file}"

    original_name = quote(filename, safe="")

    await create_test_file(filetype=file_extension, filename=filename)

    event_task_1 = asyncio.create_task(start_file_saver())

    await asyncio.sleep(0.1)

    event_task_2 = asyncio.create_task(start_file_metadata_getter())

    await asyncio.sleep(0.1)

    event_tasks = [event_task_1, event_task_2]

    content = await read_test_file(filename=filename)

    content_hex = content.hex()

    save_file_message = deepcopy(WS_MESSAGE)

    save_file_message["agent_uuid"] = file_saver_session.agent_id

    save_file_message["request_payload"] = {
        "content": content_hex,
        "filename": filename,
    }

    save_file_message["request_metadata"] = {
        "request_id": request_id,
        "session_id": session_id,
    }

    get_file_metadata_message = deepcopy(WS_MESSAGE)

    get_file_metadata_message["agent_uuid"] = metadata_session.agent_id

    try:
        async with websockets.connect(URI, additional_headers=WS_HEADERS) as websocket:
            save_file_message_json = json.dumps(save_file_message)

            await websocket.send(save_file_message_json)

            response = await websocket.recv()

            response_data = json.loads(response)

            logger.info(f"WebSocket Response: {response_data}")

            file_id = response_data.get("response", None)

            assert file_id

            get_file_metadata_message["request_payload"] = {"file_id": file_id}

            get_file_metadata_message_json = json.dumps(get_file_metadata_message)

            await websocket.send(get_file_metadata_message_json)

            response = await websocket.recv()

            response_data = json.loads(response)

            del response_data["execution_time"]

            logger.info(f"WebSocket Response: {response_data}")

            expected_response = {
                "response": {
                    "id": file_id,
                    "session_id": session_id,
                    "request_id": request_id,
                    "original_name": original_name,
                    "mimetype": mimetype,
                    "internal_id": file_id,
                    "internal_name": f"{file_id}.{file_extension}",
                    "from_agent": True,
                    "creator_id": await get_user(user_jwt_token),
                },
                "message_type": "agent_response",
            }

            assert response_data == expected_response

            logger.info("Test successful!")

    finally:
        for task in event_tasks:
            task.cancel()

        try:
            await task

        except asyncio.CancelledError:
            logger.info("Background task has been properly cancelled.")


@pytest.mark.asyncio
@pytest.mark.parametrize(
    "request_id, session_id, file_extension",
    [
        (str(uuid.uuid4()), str(uuid.uuid4()), "pdf"),
        (str(uuid.uuid4()), str(uuid.uuid4()), "txt"),
        (str(uuid.uuid4()), str(uuid.uuid4()), "xls"),
        (str(uuid.uuid4()), str(uuid.uuid4()), "docx"),
        (str(uuid.uuid4()), str(uuid.uuid4()), "jpeg"),
        (str(uuid.uuid4()), str(uuid.uuid4()), "png"),
        (str(uuid.uuid4()), str(uuid.uuid4()), "zip"),
        (str(uuid.uuid4()), str(uuid.uuid4()), "html"),
        (str(uuid.uuid4()), str(uuid.uuid4()), "mp3"),
        (str(uuid.uuid4()), str(uuid.uuid4()), "mp4"),
    ],
    ids=[
        "save pdf file and get its metadata",
        "save txt file and get its metadata",
        "save xls file and get its metadata",
        "save docs file and get its metadata",
        "save jpeg file and get its metadata",
        "save png file and get its metadata",
        "save zip file and get its metadata",
        "save html file and get its metadata",
        "save mp3 file and get its metadata",
        "save mp4 file and get its metadata",
    ],
)
async def test_agent_resave_file(
    request_id,
    session_id,
    file_extension,
    user_jwt_token: str,
    agent_jwt_factory: Callable[[str], Awaitable[str]],
    agent_factory: Callable[[], DummyAgent],
    get_user: Callable[[str], str],
):
    filesaver_jwt_token = await agent_jwt_factory(user_jwt_token)
    file_saver_session = GenAISession(jwt_token=filesaver_jwt_token)

    @file_saver_session.bind(name="File data saver", description="Saves files data")
    async def save_file_content(
        agent_context: GenAIContext, content: str, filename: str
    ):
        bytes_content = bytes.fromhex(content)

        file_id = await agent_context.files.save(
            content=bytes_content, filename=filename
        )

        return file_id

    async def start_file_saver():
        await file_saver_session.process_events()

    file_get_jwt_token = await agent_jwt_factory(user_jwt_token)
    file_session = GenAISession(jwt_token=file_get_jwt_token)

    @file_session.bind(name="File data getter", description="Gets uploaded file bytes")
    async def get_file_by_id(agent_context: GenAIContext, file_id: str):
        file = await agent_context.files.get_by_id(file_id)

        return file.getvalue().hex()

    async def start_file_getter():
        await file_session.process_events()

    file = f"test.{file_extension}"
    filename = f"{cwd}{TEST_FILES_FOLDER}/{file}"
    await create_test_file(filetype=file_extension, filename=filename)

    event_task_1 = asyncio.create_task(start_file_saver())
    await asyncio.sleep(0.1)
    event_task_2 = asyncio.create_task(start_file_getter())
    await asyncio.sleep(0.1)
    event_tasks = [event_task_1, event_task_2]
    content = await read_test_file(filename=filename)
    content_hex = content.hex()
    save_file_message = deepcopy(WS_MESSAGE)

    save_file_message["agent_uuid"] = file_saver_session.agent_id

    save_file_message["request_payload"] = {
        "content": content_hex,
        "filename": filename,
    }

    save_file_message["request_metadata"] = {
        "request_id": request_id,
        "session_id": session_id,
    }

    get_file_data_message = deepcopy(WS_MESSAGE)

    get_file_data_message["agent_uuid"] = file_session.agent_id

    try:
        async with websockets.connect(URI, additional_headers=WS_HEADERS) as websocket:
            save_file_message_json = json.dumps(save_file_message)
            await websocket.send(save_file_message_json)
            response = await websocket.recv()
            response_data = json.loads(response)
            logger.info(f"WebSocket Response: {response_data}")
            file_id = response_data.get("response", None)

            assert file_id

            get_file_data_message["request_payload"] = {"file_id": file_id}
            get_file_data_message_json = json.dumps(get_file_data_message)
            await websocket.send(get_file_data_message_json)
            response = await websocket.recv()
            response_data = json.loads(response)
            del response_data["execution_time"]
            logger.info(f"WebSocket Response: {response_data}")

            expected_response = {
                "response": content_hex,
                "message_type": "agent_response",
            }

            assert response_data == expected_response
            resave_content_hex = response_data.get("response")
            resave_content = bytes.fromhex(resave_content_hex)
            assert resave_content == content
            resave_file_message = deepcopy(WS_MESSAGE)
            resave_file_message["agent_uuid"] = file_saver_session.agent_id

            resave_file_message["request_payload"] = {
                "content": resave_content_hex,
                "filename": filename,
            }

            resave_file_message["request_metadata"] = {
                "request_id": request_id,
                "session_id": session_id,
            }

            resave_file_message_json = json.dumps(resave_file_message)
            await websocket.send(resave_file_message_json)
            resave_response = await websocket.recv()
            resave_response_data = json.loads(resave_response)
            logger.info(f"WebSocket Response: {resave_response_data}")
            resaved_file_id = resave_response_data.get("response", None)
            assert resaved_file_id
            assert file_id != resaved_file_id
            get_file_data_message["request_payload"] = {"file_id": resaved_file_id}
            get_file_data_message_json = json.dumps(get_file_data_message)
            await websocket.send(get_file_data_message_json)
            resaved_response = await websocket.recv()
            resaved_response_data = json.loads(resaved_response)
            del resaved_response_data["execution_time"]
            logger.info(f"WebSocket Response: {resaved_response_data}")

            expected_resaved_response = {
                "response": content_hex,
                "message_type": "agent_response",
            }

            assert resaved_response_data == expected_resaved_response

            logger.info("Test successful!")

    finally:
        for task in event_tasks:
            task.cancel()

        try:
            await task

        except asyncio.CancelledError:
            logger.info("Background task has been properly cancelled.")


@pytest.mark.asyncio
@pytest.mark.parametrize(
    "request_id, session_id, file_extension",
    [
        (str(uuid.uuid4()), str(uuid.uuid4()), "txt"),
        (str(uuid.uuid4()), str(uuid.uuid4()), "pdf"),
        (str(uuid.uuid4()), str(uuid.uuid4()), "xls"),
        (str(uuid.uuid4()), str(uuid.uuid4()), "docx"),
        (str(uuid.uuid4()), str(uuid.uuid4()), "jpeg"),
        (str(uuid.uuid4()), str(uuid.uuid4()), "png"),
        (str(uuid.uuid4()), str(uuid.uuid4()), "zip"),
        (str(uuid.uuid4()), str(uuid.uuid4()), "html"),
        (str(uuid.uuid4()), str(uuid.uuid4()), "mp3"),
        (str(uuid.uuid4()), str(uuid.uuid4()), "mp4"),
    ],
    ids=[
        "verify that it is possible to get txt file content",
        "verify that it is possible to get pdf file content",
        "verify that it is possible to get xls file content",
        "verify that it is possible to get docx file content",
        "verify that it is possible to get jpeg file content",
        "verify that it is possible to get png file content",
        "verify that it is possible to get zip file content",
        "verify that it is possible to get html file content",
        "verify that it is possible to get mp3 file content",
        "verify that it is possible to get mp4 file content",
    ],
)
async def test_agent_download_file(
    request_id,
    session_id,
    file_extension,
    user_jwt_token: str,
    agent_jwt_factory: Callable[[str], Awaitable[str]],
):
    filesaver_jwt_token = await agent_jwt_factory(user_jwt_token)
    file_saver_session = GenAISession(jwt_token=filesaver_jwt_token)

    @file_saver_session.bind(name="File data saver", description="Saves files data")
    async def save_file_content(
        agent_context: GenAIContext, content: str, filename: str
    ):
        bytes_content = bytes.fromhex(content)

        file_id = await agent_context.files.save(
            content=bytes_content, filename=filename
        )

        return file_id

    async def start_file_saver():
        await file_saver_session.process_events()

    file_get_jwt_token = await agent_jwt_factory(user_jwt_token)
    file_session = GenAISession(jwt_token=file_get_jwt_token)

    @file_session.bind(name="File data getter", description="Gets uploaded file bytes")
    async def get_file_by_id(agent_context: GenAIContext, file_id: str):
        file = await agent_context.files.get_by_id(file_id)

        return file.getvalue().hex()

    async def start_file_getter():
        await file_session.process_events()

    file = f"test.{file_extension}"
    filename = f"{cwd}{TEST_FILES_FOLDER}/{file}"
    await create_test_file(filetype=file_extension, filename=filename)

    event_task_1 = asyncio.create_task(start_file_saver())
    await asyncio.sleep(0.1)
    event_task_2 = asyncio.create_task(start_file_getter())
    await asyncio.sleep(0.1)

    event_tasks = [event_task_1, event_task_2]
    uploaded_file_content = await read_test_file(filename=filename)
    content_hex = uploaded_file_content.hex()
    save_file_message = deepcopy(WS_MESSAGE)
    save_file_message["agent_uuid"] = file_saver_session.agent_id

    save_file_message["request_payload"] = {
        "content": content_hex,
        "filename": filename,
    }

    save_file_message["request_metadata"] = {
        "request_id": request_id,
        "session_id": session_id,
    }

    get_file_message = deepcopy(WS_MESSAGE)

    get_file_message["agent_uuid"] = file_session.agent_id

    try:
        async with websockets.connect(URI, additional_headers=WS_HEADERS) as websocket:
            save_file_message_json = json.dumps(save_file_message)

            await websocket.send(save_file_message_json)
            response = await websocket.recv()
            response_data = json.loads(response)
            logger.info(f"WebSocket Response: {response_data}")
            file_id = response_data.get("response", None)

            assert file_id

            get_file_message["request_payload"] = {"file_id": file_id}
            get_file_message_json = json.dumps(get_file_message)
            await websocket.send(get_file_message_json)
            response = await websocket.recv()
            response_data = json.loads(response)
            del response_data["execution_time"]

            logger.info(f"WebSocket Response: {response_data}")
            downloaded_file_content = bytes.fromhex(response_data.get("response"))
            assert uploaded_file_content == downloaded_file_content
            logger.info("Test successful!")

    finally:
        for task in event_tasks:
            task.cancel()

        try:
            await task

        except asyncio.CancelledError:
            logger.info("Background task has been properly cancelled.")


@pytest.mark.asyncio
@pytest.mark.parametrize(
    "request_id, session_id, file_extension, content",
    [
        (str(uuid.uuid4()), str(uuid.uuid4()), "txt", None),
        (str(uuid.uuid4()), str(uuid.uuid4()), "txt", 42),
        (str(uuid.uuid4()), str(uuid.uuid4()), "txt", 3.14),
        (str(uuid.uuid4()), str(uuid.uuid4()), "txt", 2 + 3j),
        (str(uuid.uuid4()), str(uuid.uuid4()), "txt", True),
        (str(uuid.uuid4()), str(uuid.uuid4()), "txt", "hello"),
        (str(uuid.uuid4()), str(uuid.uuid4()), "txt", [1, 2, 3]),
        (str(uuid.uuid4()), str(uuid.uuid4()), "txt", (1, 2, 3)),
        (str(uuid.uuid4()), str(uuid.uuid4()), "txt", range(5)),
        (str(uuid.uuid4()), str(uuid.uuid4()), "txt", {"key": "value"}),
        (str(uuid.uuid4()), str(uuid.uuid4()), "txt", {1, 2, 3}),
        (str(uuid.uuid4()), str(uuid.uuid4()), "txt", frozenset([1, 2])),
        (
            str(uuid.uuid4()),
            str(uuid.uuid4()),
            "txt",
            bytearray(b"hi"),
        ),  # TODO: Is it allowed?
        (
            str(uuid.uuid4()),
            str(uuid.uuid4()),
            "txt",
            memoryview(b"hi"),
        ),  # TODO: Is it allowed?
    ],
    ids=[
        "save txt file with None",
        "save txt file with int",
        "save txt file with float",
        "save txt file with complex number",
        "save txt file with bool",
        "save txt file with string",
        "save txt file with list",
        "save txt file with tuple",
        "save txt file with range",
        "save txt file with dict",
        "save txt file with set",
        "save txt file with frozenset",
        "save txt file with bytearray",
        "save txt file with memoryview",
    ],
)
async def test_agent_invalid_save_file(
    request_id,
    session_id,
    file_extension,
    content,
    user_jwt_token: str,
    agent_jwt_factory: Callable[[str], Awaitable[str]],
    agent_factory: Callable[[], DummyAgent],
):
    filesaver_jwt_token = await agent_jwt_factory(user_jwt_token)
    file_saver_session = GenAISession(jwt_token=filesaver_jwt_token)

    filename = f"{cwd}{TEST_FILES_FOLDER}/test.{file_extension}"

    @file_saver_session.bind(name="File data saver", description="Saves files data")
    async def save_file_content(agent_context: GenAIContext):
        logger.info(f"Content: {content}")
        file_id = await agent_context.files.save(content=content, filename=filename)
        return file_id

    async def start_file_saver():
        await file_saver_session.process_events()

    metadata_jwt_token = await agent_jwt_factory(user_jwt_token)
    metadata_session = GenAISession(jwt_token=metadata_jwt_token)

    @metadata_session.bind(
        name="File metadata getter", description="Gets metadata of uploaded file"
    )
    async def get_file_metadata_by_id(agent_context: GenAIContext, file_id: str):
        metadata = await agent_context.files.get_metadata_by_id(file_id)

        return metadata

    async def start_file_metadata_getter():
        await metadata_session.process_events()

    event_task_1 = asyncio.create_task(start_file_saver())
    await asyncio.sleep(0.1)
    event_task_2 = asyncio.create_task(start_file_metadata_getter())
    await asyncio.sleep(0.1)

    event_tasks = [event_task_1, event_task_2]
    save_file_message = deepcopy(WS_MESSAGE)
    save_file_message["agent_uuid"] = file_saver_session.agent_id
    save_file_message["request_metadata"] = {
        "request_id": request_id,
        "session_id": session_id,
    }

    get_file_metadata_message = deepcopy(WS_MESSAGE)

    get_file_metadata_message["agent_uuid"] = metadata_session.agent_id

    try:
        async with websockets.connect(URI, additional_headers=WS_HEADERS) as websocket:
            save_file_message_json = json.dumps(save_file_message)
            await websocket.send(save_file_message_json)
            response = await websocket.recv()
            response_data = json.loads(response)
            logger.info(f"WebSocket Response: {response_data}")

            error = response_data.get("error")
            error_message = error.get("error_message")
            execution_time = response_data.get("execution_time")
            message_type = response_data.get("message_type")

            assert error
            assert error_message  # TODO: Add error message processing
            assert execution_time == 0
            assert message_type == "agent_error"

            logger.info("Test successful!")

    finally:
        for task in event_tasks:
            task.cancel()

        try:
            await task

        except asyncio.CancelledError:
            logger.info("Background task has been properly cancelled.")
